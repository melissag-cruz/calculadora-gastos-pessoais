"""Exporta o resumo de um mês para Excel e gera um gráfico de gastos por categoria.

As bibliotecas externas (openpyxl e matplotlib) só são importadas quando usadas, para que o
restante do programa funcione mesmo sem elas instaladas.
"""

from datetime import date
from decimal import Decimal
from pathlib import Path

from gastos import ResumoMensal, formatar_reais

PASTA_RELATORIOS = Path(__file__).with_name("relatorios")
FORMATO_MOEDA = '"R$" #,##0.00'


def exportar_excel(resumo: ResumoMensal, destino: Path | None = None) -> Path:
    """Cria uma planilha com as despesas do mês e uma aba de resumo. Devolve o caminho salvo."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    destino = destino or PASTA_RELATORIOS / f"relatorio-{resumo.mes}.xlsx"
    destino.parent.mkdir(parents=True, exist_ok=True)

    planilha = Workbook()
    aba_despesas = planilha.active
    aba_despesas.title = "Despesas"
    aba_despesas.append(["Data", "Descrição", "Categoria", "Valor"])
    for despesa in resumo.despesas:
        aba_despesas.append(
            [
                date.fromisoformat(despesa["data"]),
                despesa["descricao"],
                despesa["categoria"],
                Decimal(despesa["valor"]),
            ]
        )

    aba_resumo = planilha.create_sheet("Resumo")
    aba_resumo.append(["Categoria", "Gasto", "Orçamento", "Passou do orçamento"])
    for categoria, total in sorted(resumo.totais_por_categoria.items()):
        excesso = resumo.excesso(categoria)
        aba_resumo.append([categoria, total, resumo.orcamentos.get(categoria), excesso or None])
    aba_resumo.append([])
    aba_resumo.append(["Total", resumo.total])
    if resumo.renda is not None:
        aba_resumo.append(["Renda", resumo.renda])
        aba_resumo.append(["Saldo", resumo.renda - resumo.total])

    for aba, larguras in [(aba_despesas, [12, 30, 18, 14]), (aba_resumo, [18, 14, 14, 20])]:
        for celula in aba[1]:
            celula.font = Font(bold=True)
        for coluna, largura in zip("ABCD", larguras, strict=True):
            aba.column_dimensions[coluna].width = largura
        for linha in aba.iter_rows(min_row=2):
            for celula in linha:
                if isinstance(celula.value, Decimal):
                    celula.number_format = FORMATO_MOEDA
                elif isinstance(celula.value, date):
                    celula.number_format = "DD/MM/YYYY"
    for linha in aba_resumo.iter_rows(min_row=len(resumo.totais_por_categoria) + 3):
        linha[0].font = Font(bold=True)

    planilha.save(destino)
    return destino


def gerar_grafico(resumo: ResumoMensal, destino: Path | None = None) -> Path:
    """Salva um gráfico de barras com o gasto de cada categoria no mês. Devolve o caminho salvo.

    Categorias que passaram do orçamento aparecem em vermelho, e o limite é marcado com uma
    linha tracejada.
    """
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    destino = destino or PASTA_RELATORIOS / f"grafico-{resumo.mes}.png"
    destino.parent.mkdir(parents=True, exist_ok=True)

    totais = resumo.totais_por_categoria
    categorias = sorted(totais, key=lambda categoria: totais[categoria], reverse=True)
    cores = ["#d9534f" if resumo.excesso(c) else "#4a90d9" for c in categorias]

    figura, eixo = plt.subplots(figsize=(8, 1.5 + 0.6 * len(categorias)))
    barras = eixo.barh(categorias, [float(totais[c]) for c in categorias], color=cores)
    eixo.invert_yaxis()
    eixo.bar_label(barras, labels=[formatar_reais(totais[c]) for c in categorias], padding=4)

    rotulo = "Orçamento"
    for posicao, categoria in enumerate(categorias):
        limite = resumo.orcamentos.get(categoria)
        if limite is not None:
            eixo.vlines(
                float(limite),
                posicao - 0.4,
                posicao + 0.4,
                colors="black",
                linestyles="--",
                label=rotulo,
            )
            rotulo = "_nolegend_"
    if rotulo == "_nolegend_":
        eixo.legend(loc="lower right")

    eixo.set_title(f"Gastos por categoria em {resumo.mes} (total {formatar_reais(resumo.total)})")
    eixo.set_xlabel("Valor (R$)")
    eixo.margins(x=0.25)
    eixo.spines[["top", "right"]].set_visible(False)
    figura.tight_layout()
    figura.savefig(destino, dpi=120)
    plt.close(figura)
    return destino
