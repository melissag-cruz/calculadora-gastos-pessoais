"""Testes da exportação para Excel e do gráfico (precisam de openpyxl e matplotlib)."""

import importlib.util
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

import relatorios
from gastos import ResumoMensal, criar_despesa

TEM_OPENPYXL = importlib.util.find_spec("openpyxl") is not None
TEM_MATPLOTLIB = importlib.util.find_spec("matplotlib") is not None


def resumo_de_exemplo(renda: Decimal | None = Decimal("3000")) -> ResumoMensal:
    despesas = [
        criar_despesa("2026-10-05", "Mercado", "Alimentação", Decimal("485.50")),
        criar_despesa("2026-10-10", "Ônibus", "Transporte", Decimal("4.40")),
        criar_despesa("2026-10-12", "Restaurante", "Alimentação", Decimal("62.90")),
    ]
    return ResumoMensal(
        mes="2026-10",
        despesas=despesas,
        totais_por_categoria={"Alimentação": Decimal("548.40"), "Transporte": Decimal("4.40")},
        total=Decimal("552.80"),
        renda=renda,
        orcamentos={"Alimentação": Decimal("500.00")},
    )


class TesteComPasta(unittest.TestCase):
    def setUp(self):
        pasta = tempfile.TemporaryDirectory()
        self.addCleanup(pasta.cleanup)
        self.pasta = Path(pasta.name)


@unittest.skipUnless(TEM_OPENPYXL, "openpyxl não instalado")
class TestExportarExcel(TesteComPasta):
    def test_cria_abas_de_despesas_e_resumo(self):
        from openpyxl import load_workbook

        caminho = relatorios.exportar_excel(resumo_de_exemplo(), self.pasta / "relatorio.xlsx")
        planilha = load_workbook(caminho)
        self.assertEqual(planilha.sheetnames, ["Despesas", "Resumo"])

        despesas = list(planilha["Despesas"].iter_rows(values_only=True))
        self.assertEqual(despesas[0], ("Data", "Descrição", "Categoria", "Valor"))
        self.assertEqual(len(despesas), 4)
        self.assertEqual(despesas[1][1:], ("Mercado", "Alimentação", 485.5))

        resumo = {linha[0]: linha[1:] for linha in planilha["Resumo"].iter_rows(values_only=True)}
        self.assertEqual(resumo["Alimentação"], (548.4, 500, 48.4))
        self.assertEqual(resumo["Total"][0], 552.8)
        self.assertEqual(resumo["Saldo"][0], 2447.2)

    def test_sem_renda_nao_tem_saldo(self):
        from openpyxl import load_workbook

        caminho = relatorios.exportar_excel(resumo_de_exemplo(renda=None), self.pasta / "r.xlsx")
        rotulos = [
            linha[0] for linha in load_workbook(caminho)["Resumo"].iter_rows(values_only=True)
        ]
        self.assertNotIn("Saldo", rotulos)


@unittest.skipUnless(TEM_MATPLOTLIB, "matplotlib não instalado")
class TestGerarGrafico(TesteComPasta):
    def test_salva_imagem_png(self):
        caminho = relatorios.gerar_grafico(resumo_de_exemplo(), self.pasta / "grafico.png")
        self.assertEqual(caminho.read_bytes()[:8], b"\x89PNG\r\n\x1a\n")


if __name__ == "__main__":
    unittest.main()
